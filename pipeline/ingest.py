"""Extract plain text from an uploaded file.

Supported on Day 2: PDF, DOCX, EML, plain text / markdown.
.doc and .msg are explicitly not supported in v0.1 — they need LibreOffice
or extra Windows libs respectively. Caller gets a clear error.
"""

from __future__ import annotations

from pathlib import Path

import pdfplumber
from docx import Document

try:
    import mailparser
except ImportError:  # pragma: no cover
    mailparser = None


class UnsupportedFormat(Exception):
    pass


def detect_kind(filename: str) -> str:
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext == "pdf":
        return "pdf"
    if ext == "docx":
        return "docx"
    if ext == "eml":
        return "eml"
    if ext in {"txt", "md"}:
        return "text"
    return ext or "unknown"


def extract_text(path: str | Path) -> str:
    path = Path(path)
    kind = detect_kind(path.name)

    if kind == "pdf":
        out: list[str] = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                out.append(page.extract_text() or "")
        return "\n\n".join(out).strip()

    if kind == "docx":
        doc = Document(str(path))
        chunks: list[str] = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(c.text for c in row.cells))
        return "\n".join(chunks).strip()

    if kind == "eml":
        if mailparser is None:
            raise UnsupportedFormat("mailparser not installed")
        mail = mailparser.parse_from_file(str(path))
        body_parts = mail.text_plain or []
        if not body_parts and mail.body:
            body_parts = [mail.body]
        header = (
            f"From: {mail.from_}\n"
            f"To: {mail.to}\n"
            f"Date: {mail.date}\n"
            f"Subject: {mail.subject}\n\n"
        )
        return (header + "\n\n".join(body_parts)).strip()

    if kind == "text":
        return path.read_text(errors="replace").strip()

    raise UnsupportedFormat(f"Unsupported file type: .{kind}")
